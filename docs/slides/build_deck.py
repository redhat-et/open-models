#!/usr/bin/env python3
"""Build the OCTO Open Models lightning deck (6 slides, ~5 min).

Run: python3 build_deck.py (needs python-pptx); rewrites the .pptx next to itself.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PURPLE = RGBColor(0x4F, 0x1F, 0x8F)
BLUE = RGBColor(0x0F, 0x3B, 0x8C)
TEAL = RGBColor(0x0B, 0x6A, 0x7D)
GREEN = RGBColor(0x1E, 0xC4, 0x57)
DARK = RGBColor(0x15, 0x15, 0x15)
GREY = RGBColor(0x6A, 0x6E, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def bar(slide, color=PURPLE, y=0.0, h=0.12):
    s = slide.shapes.add_shape(1, 0, Inches(y), prs.slide_width, Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()
    return s


def text(slide, x, y, w, h, s, size, color=DARK, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = s
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb


def bullets(slide, x, y, w, items, size=20, gap_after=10):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(5.6))
    tf = tb.text_frame; tf.word_wrap = True
    for i, (s, sub) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap_after)
        r = p.add_run(); r.text = "▸ " if not sub else "   "
        r.font.size = Pt(size); r.font.color.rgb = TEAL
        r2 = p.add_run(); r2.text = s
        r2.font.size = Pt(size if not sub else size - 4)
        r2.font.bold = not sub
        r2.font.color.rgb = DARK if not sub else GREY
    return tb


# 1 - title
s = prs.slides.add_slide(BLANK); bg(s, DARK)
bar(s, PURPLE, 0, 0.18)
text(s, 0.9, 2.2, 11.5, 1.4, "🐙 OCTO Open Models", 54, WHITE, True)
text(s, 0.9, 3.4, 11.5, 0.9, "Open models, one key, zero invoice", 28, RGBColor(0x73, 0xBC, 0xF7))
text(s, 0.9, 6.2, 11.5, 0.6, "Launching 2026-09-09  ·  octo-open-models  ·  #octo-open-model-platform", 16, GREY)

# 2 - problem
s = prs.slides.add_slide(BLANK)
bar(s, BLUE)
text(s, 0.9, 0.5, 11.5, 0.8, "Why this exists", 34, PURPLE, True)
bullets(s, 0.9, 1.6, 11.5, [
    ("Token spend is a real line item", None),
    ("Associates hit monthly budgets and stop shipping", True),
    ("…or spin up shadow GPU spend nobody tracks", True),
    ("Open models are finally good enough for daily agentic work", None),
    ("…but serving them well (quant, spec-decode, context, concurrency) is a project", True),
    ("…one every team would otherwise repeat badly on its own", True),
], size=22)

# 3 - what we built
s = prs.slides.add_slide(BLANK)
bar(s, TEAL)
text(s, 0.9, 0.5, 11.5, 0.8, "What we built: four ones", 34, PURPLE, True)
bullets(s, 0.9, 1.6, 11.5, [
    ("One gateway — LiteLLM on OpenShift", None),
    ("OpenAI-compatible, per-user keys + budgets, auto fallback, spend + latency metrics", True),
    ("One portal — GitHub login, self-serve keys, live model cards, savings counter", None),
    ("One source of truth — catalog.yaml", None),
    ("Gateway config, dashboards, alerts, user configs: all generated. Onboarding a model = a PR", True),
    ("One command for users", None),
    ("Paste a curl → opencode pointed at the gateway. ~2 minutes, nothing to install", True),
], size=22)

# 4 - tiers
s = prs.slides.add_slide(BLANK)
bar(s, GREEN)
text(s, 0.9, 0.5, 11.5, 0.8, "Tiers: roles, not speed", 34, PURPLE, True)
rows = [
    ("Tier", "Model", "Hardware", "For", "$/Mtok in/out"),
    ("🏆 frontier", "Qwen 3.8 27B FP8, 131k ctx", "rented MI300X 192GB", "hardest agentic work; falls back to core", "0.60 / 2.40"),
    ("⚡ core", "Flash-Next NVFP4, 262k ctx", "4x L40S", "daily driver · your default · 10 seats @ 19 tok/s", "0.20 / 1.00"),
    ("📦 bulk", "27B NVFP4, 32k ctx", "RTX 5090", "batch + background · 8 seats @ 12.5 tok/s", "0.30 / 1.20"),
    ("🧪 lab", "27B FP8 / Flash-Next", "DGX Sparks", "bake-offs, VPN endpoints, no SLA", "unmetered"),
]
tbl = s.shapes.add_table(5, 5, Inches(0.7), Inches(1.55), Inches(12), Inches(4.6)).table
for i, w in enumerate((1.7, 3.4, 2.2, 3.6, 1.7)):
    tbl.columns[i].width = Inches(w)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = val
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(14 if r else 15)
        p.font.bold = (r == 0) or (c == 0)
        p.font.color.rgb = WHITE if r == 0 else DARK
        cell.fill.solid()
        cell.fill.fore_color.rgb = PURPLE if r == 0 else (RGBColor(0xF5, 0xF5, 0xF5) if r % 2 else WHITE)
text(s, 0.9, 6.5, 11.5, 0.5, "Every seat number measured on the live box, 2026-09-08. Notional rates vs frontier APIs.", 14, GREY)

# 5 - money
s = prs.slides.add_slide(BLANK)
bar(s, BLUE)
text(s, 0.9, 0.5, 11.5, 0.8, "We watch money like users do", 34, PURPLE, True)
bullets(s, 0.9, 1.6, 11.5, [
    ("Savings = real numbers, not vibes", None),
    ("frontier-API x8 notional, per user + per model, live in Grafana", True),
    ("Own GPUs sleep when we do", None),
    ("L40S node torn down 18:00, back 04:00 weekdays", True),
    ("Rented compute: delete-not-stop policy", None),
    ("meter only stops when the VM is gone; wake/provision alerts Slack", True),
    ("If a model can't earn its $/hr, it goes", None),
    ("Retired GLM-4.5-Air the day its tool-calling failed our own usage — same-day teardown, no sunk-cost drift", True),
], size=22)

# 6 - CTA
s = prs.slides.add_slide(BLANK); bg(s, DARK)
bar(s, PURPLE, 0, 0.18)
text(s, 0.9, 1.0, 11.5, 0.9, "Log in → create a key → paste one line → ship", 36, WHITE, True)
tb = s.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.1))
p = tb.text_frame.paragraphs[0].add_run()
p.text = "curl -fsSL <portal>/setup.sh | OCTO_KEY=•••• bash"
p.font.size = Pt(22); p.font.color.rgb = RGBColor(0x73, 0xBC, 0xF7)
bullets(s, 0.9, 3.9, 11.5, [
    ("Portal: octo-app-octo-models.apps.emerg.pcbk.p1.openshiftapps.com", None),
    ("Live dashboard needs no login — links off the portal home", True),
    ("Repo: github.com/redhat-et/open-models — model onboarding is a PR", True),
    ("Feedback: #octo-open-model-platform", True),
    ("Try it during this demo: the savings counter updates live", None),
], size=20, gap_after=8)

prs.save(str(Path(__file__).with_name("octo-open-models-launch.pptx")))
print("saved", len(prs.slides._sldIdLst), "slides")
